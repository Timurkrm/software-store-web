from pathlib import Path
from docx import Document
from pypdf import PdfReader
import json, hashlib

root=Path(r'C:\Users\timur\Downloads\software-store-web\software-store-web')
out=Path('F:/SoftStore-course-document')
ref=Path(r'C:\Users\timur\Downloads\Пояснительная записка(1).docx')
doc=Document(ref)
text=[]
for p in doc.paragraphs:
    if p.text: text.append(f'[{p.style.style_id}] {p.text}')
for i,t in enumerate(doc.tables):
    text.append(f'TABLE {i}')
    for row in t.rows: text.append(' | '.join(c.text for c in row.cells))
(out/'reference.txt').write_text('\n'.join(text),encoding='utf-8')
spec={'sha256':hashlib.sha256(ref.read_bytes()).hexdigest(),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'images':len(doc.inline_shapes),'sections':[],'styles':[]}
for s in doc.sections:
    spec['sections'].append({k:getattr(s,k).cm if getattr(s,k) is not None else None for k in ['page_width','page_height','left_margin','right_margin','top_margin','bottom_margin','header_distance','footer_distance']})
for s in doc.styles:
    if s.type==1:
        f=s.paragraph_format
        spec['styles'].append({'id':s.style_id,'name':s.name,'font':s.font.name,'size':s.font.size.pt if s.font.size else None,'alignment':str(f.alignment),'first_indent':f.first_line_indent.cm if f.first_line_indent else None,'line_spacing':str(f.line_spacing)})
(out/'reference-spec.json').write_text(json.dumps(spec,ensure_ascii=False,indent=2),encoding='utf-8')
pdf=next(root.glob('МУ*.pdf'))
reader=PdfReader(pdf)
(out/'requirements.txt').write_text('\n'.join(f'\n=== PAGE {i+1} ===\n{p.extract_text()}' for i,p in enumerate(reader.pages)),encoding='utf-8')
print('Requirements pages',len(reader.pages),'reference',spec['paragraphs'],'paragraphs',spec['tables'],'tables',spec['images'],'images')
